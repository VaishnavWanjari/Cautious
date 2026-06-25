"""Tests for the Excel/PDF/HTML exporters.

The xlsx/pdf exporters need openpyxl/reportlab; those tests are skipped if the
optional dependency is missing (e.g. an offline dev box) but run fully in CI.
The HTML exporter is pure stdlib and always runs.
"""

import pytest

_SCHEDULE = {
    "activities": [
        {
            "activity_id": "A-010",
            "name": "Hydrotest",
            "duration": 7,
            "start_date": "2026-12-21",
            "finish_date": "2026-12-29",
            "es": 0,
            "ef": 6,
            "ls": 0,
            "lf": 6,
            "total_float": 0,
            "is_critical": True,
        },
        {
            "activity_id": "A-020",
            "name": "Dewatering",
            "duration": 2,
            "start_date": "2026-12-30",
            "finish_date": "2026-12-31",
            "es": 7,
            "ef": 8,
            "ls": 7,
            "lf": 8,
            "total_float": 0,
            "is_critical": True,
        },
    ]
}

_SUMMARY = {
    "project": {"name": "Train 1", "client": "ACME", "mechanical_completion_date": "2027-01-15"},
    "total_activities": 2,
    "completed": 0,
    "in_progress": 1,
    "pending": 1,
    "critical": 2,
}


def test_html_export_contains_activities():
    from app.infrastructure.exporters.html_exporter import export_dashboard_html

    html = export_dashboard_html(_SUMMARY, _SCHEDULE)
    assert "Hydrotest" in html
    assert "Train 1" in html
    assert "crit" in html  # critical row styling applied


def test_xlsx_export_is_valid_workbook():
    pytest.importorskip("openpyxl")
    import io

    from openpyxl import load_workbook

    from app.infrastructure.exporters.xlsx_exporter import export_schedule_xlsx

    data = export_schedule_xlsx("Train 1", _SCHEDULE)
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Schedule"]
    # title + blank + header + 2 data rows
    values = [c.value for c in ws["B"] if c.value]
    assert "Hydrotest" in values


def test_pdf_export_produces_bytes():
    pytest.importorskip("reportlab")

    from app.infrastructure.exporters.pdf_exporter import export_schedule_pdf

    data = export_schedule_pdf("Train 1", _SCHEDULE)
    assert data[:4] == b"%PDF"
