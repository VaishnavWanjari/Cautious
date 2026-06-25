"""Excel (.xlsx) schedule export via OpenPyXL."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_CRIT_FILL = PatternFill("solid", fgColor="F8CBAD")

_COLUMNS = [
    ("Activity ID", "activity_id"),
    ("Activity Name", "name"),
    ("Duration (d)", "duration"),
    ("Start", "start_date"),
    ("Finish", "finish_date"),
    ("ES", "es"),
    ("EF", "ef"),
    ("LS", "ls"),
    ("LF", "lf"),
    ("Total Float", "total_float"),
    ("Critical", "is_critical"),
]


def export_schedule_xlsx(project_name: str, schedule: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"

    ws["A1"] = f"Commissioning Schedule — {project_name}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])

    header_row = 3
    for col, (label, _key) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for r, act in enumerate(schedule.get("activities", []), start=header_row + 1):
        for col, (_label, key) in enumerate(_COLUMNS, start=1):
            val = act.get(key)
            if key == "is_critical":
                val = "Yes" if val else "No"
            elif key in ("start_date", "finish_date") and val is not None:
                val = str(val)
            cell = ws.cell(row=r, column=col, value=val)
            if act.get("is_critical"):
                cell.fill = _CRIT_FILL

    for col in range(1, len(_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
