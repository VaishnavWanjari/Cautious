"""Export module data to Excel / CSV / JSON.

Uses the same column specs as the tables so exports match what the user sees,
including computed columns (balance, variance, …).
"""

from __future__ import annotations

import csv
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from . import repository
from .columns import spec_for


def _table(entity_key: str) -> tuple[list[str], list[list[Any]]]:
    specs = spec_for(entity_key)
    headers = [s.label for s in specs]
    rows: list[list[Any]] = []
    for obj in repository.list_all(entity_key):
        rows.append([_cell(getattr(obj, s.name, "")) for s in specs])
    return headers, rows


def _cell(v: Any) -> Any:
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    return v


def export_excel(entity_key: str, path: str | Path, title: str | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    headers, rows = _table(entity_key)
    wb = Workbook()
    ws = wb.active
    ws.title = (title or entity_key)[:31]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, min(48, len(str(h)) + 6))
    ws.freeze_panes = "A2"
    path = Path(path)
    wb.save(path)
    return path


def export_csv(entity_key: str, path: str | Path) -> Path:
    headers, rows = _table(entity_key)
    path = Path(path)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)
    return path


def export_json(entity_key: str, path: str | Path) -> Path:
    specs = spec_for(entity_key)
    out = []
    for obj in repository.list_all(entity_key):
        out.append({s.name: _cell(getattr(obj, s.name, "")) for s in specs})
    path = Path(path)
    path.write_text(json.dumps(out, indent=2, default=str), "utf-8")
    return path
