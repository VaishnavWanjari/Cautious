"""Generate sample Excel import templates into ``desktop/sample``.

These match the app's import headers (the field labels), so users can see the
expected layout for bulk-importing each module. Run: ``python tools/make_samples.py``.
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.core.columns import spec_for  # noqa: E402

SAMPLE_DIR = Path(__file__).resolve().parent.parent / "sample"

# A couple of example rows per module (by field name) to show the format.
EXAMPLES = {
    "pmcc": [
        {"number": "PMCC-01", "description": "Substation & EDG handover",
         "system": "Electrical", "area": "Utilities", "discipline": "Electrical",
         "status": "In Progress", "priority": "High", "completion": 45,
         "target_date": "2027-04-26"},
        {"number": "PMCC-02", "description": "Inlet gas manifold",
         "system": "Feed Gas", "area": "Inlet Facility", "discipline": "Piping",
         "status": "Not Started", "priority": "Critical", "completion": 0},
    ],
    "equipment": [
        {"tag_number": "P-3001", "description": "Booster pump", "area": "Gas Train-3",
         "discipline": "Rotary Equipment", "equipment_type": "Rotary",
         "status": "Installed", "completion": 20},
    ],
    "punch": [
        {"number": "PL-1001", "description": "Missing gasket", "category": "B",
         "area": "AGRU", "discipline": "Piping", "priority": "Medium",
         "status": "Open", "raised_date": "2026-06-01"},
    ],
    "consumable": [
        {"name": "Nitrogen", "unit": "Nm3", "required": 5000, "available": 3200,
         "consumed": 1800},
    ],
}


def main() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)
    for key in ("pmcc", "equipment", "punch", "consumable", "procedure",
                "preservation", "manpower", "special_activity", "proposal_item"):
        specs = [s for s in spec_for(key) if not s.read_only]
        wb = Workbook()
        ws = wb.active
        ws.title = key[:31]
        ws.append([s.label for s in specs])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in EXAMPLES.get(key, []):
            ws.append([row.get(s.name, "") for s in specs])
        out = SAMPLE_DIR / f"sample_{key}.xlsx"
        wb.save(out)
        print("wrote", out)


if __name__ == "__main__":
    main()
